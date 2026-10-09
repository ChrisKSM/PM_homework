#!/usr/bin/env python3
"""모델 검증 일정 상세 — Excel(openpyxl) → ModelRow JSON.

시트 레이아웃 (첨부 엑셀 / export 스크립트 출력과 동일):
  - 메타 열: 제품군, 모델명, 이벤트, 개발등급, 생산(업체), SoC, 담당, 구분, 주요 변경점, Status
  - 우측: 일별 Gantt (날짜 헤더 1~2행)

Usage:
  pip install openpyxl
  python3 scripts/import-verification-schedule-from-xlsx.py \\
    scripts/data/verification-schedule-detail.xlsx \\
    -o scripts/seed-model-schedule-data.json
"""
from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.cell.cell import Cell
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError as e:
    raise SystemExit("openpyxl 필요: pip install openpyxl") from e

TEST_TYPES = ("일반성능", "호환성", "안정성", "시너지")
STATUS_VALUES = ("완료", "검증제외", "예정", "진행중", "NG", "지연")

META_ALIASES: dict[str, tuple[str, ...]] = {
    "category": ("제품군", "카테고리", "product"),
    "model": ("모델명", "model"),
    "event": ("이벤트", "event"),
    "variant": ("개발등급", "등급"),
    "manufacturer": ("생산업체", "생산", "manufacturer"),
    "soc": ("soc", "SoC"),
    "staff": ("담당", "담당스탭", "담당스텝", "staff"),
    "testType": ("구분", "test"),
    "changes": ("주요 변경", "주요 변경점", "변경"),
    "status": ("status", "Status", "상태"),
}


def norm_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).replace("\r\n", "\n").strip()


def slug_id(*parts: str) -> str:
    raw = "-".join(p for p in parts if p)
    s = re.sub(r"[^a-zA-Z0-9가-힣]+", "-", raw).strip("-").lower()
    return s[:80] or "row"


def norm_status(raw: str) -> str:
    s = norm_text(raw)
    if not s:
        return "예정"
    for st in STATUS_VALUES:
        if st in s or s == st:
            return st
    if "완" in s:
        return "완료"
    if "제외" in s:
        return "검증제외"
    if "NG" in s.upper():
        return "NG"
    if "진행" in s:
        return "진행중"
    if "지연" in s:
        return "지연"
    return "예정"


def norm_test_type(raw: str) -> str:
    s = norm_text(raw)
    for tt in TEST_TYPES:
        if tt in s:
            return tt
    return TEST_TYPES[0]


def cell_rgb(cell: Cell) -> str | None:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return None
    c = fill.fgColor
    if c is None:
        return None
    if c.type == "rgb" and c.rgb:
        rgb = c.rgb
        if len(rgb) == 8:
            rgb = rgb[2:]
        return rgb.upper()
    return None


def bar_type_from_rgb(rgb: str | None) -> str:
    if not rgb:
        return "planned"
    r = rgb.upper()
    # Excel / UI 팔레트 근사
    if r in ("EF4444", "FF0000", "C00000", "E74C3C", "FF6B6B"):
        return "event_ng"
    if r in ("22C55E", "00B050", "92D050", "70AD47"):
        return "event_ok"
    if r in ("FACC15", "FFC000", "FFFF00", "FFD966", "FBBF24"):
        return "inprogress"
    if r in ("A78BFA", "7030A0", "9966FF", "B4A7D6"):
        return "event_done_est"
    if r in ("F97316", "ED7D31", "FF9900", "FFA500"):
        return "su_fota"
    if r in ("94A3B8", "4472C4", "5B9BD5", "0070C0", "0000FF"):
        return "planned"
    # 밝기 기반 fallback
    try:
        rr, gg, bb = int(r[0:2], 16), int(r[2:4], 16), int(r[4:6], 16)
    except Exception:
        return "planned"
    if rr > 200 and gg < 120 and bb < 120:
        return "event_ng"
    if gg > 180 and rr < 160:
        return "event_ok"
    if rr > 200 and gg > 180:
        return "inprogress"
    return "planned"


def parse_date_header(v: Any, default_year: int) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = norm_text(v)
    if not s:
        return None
    m = re.match(r"^(\d{1,2})[/\.-](\d{1,2})$", s)
    if m:
        mo, da = int(m.group(1)), int(m.group(2))
        y = default_year
        if mo >= 9:
            y = default_year
        else:
            y = default_year + 1
        try:
            return date(y, mo, da)
        except ValueError:
            return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    return None


def find_meta_header_row(ws: Worksheet) -> tuple[int, dict[str, int]]:
    for r in range(1, min(ws.max_row, 40) + 1):
        row_vals = {c: norm_text(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 80) + 1)}
        texts = " ".join(row_vals.values())
        if "모델" in texts and "구분" in texts:
            cols: dict[str, int] = {}
            for c, val in row_vals.items():
                low = val.lower()
                for key, aliases in META_ALIASES.items():
                    if key in cols:
                        continue
                    if any(a.lower() in low or low == a.lower() for a in aliases if a):
                        cols[key] = c
            if "model" in cols and "testType" in cols:
                return r, cols
    raise ValueError("헤더 행을 찾을 수 없습니다 (모델명·구분 열 확인)")


def find_timeline(ws: Worksheet, header_row: int, meta_cols: dict[str, int]) -> tuple[list[date], int]:
    meta_max = max(meta_cols.values()) if meta_cols else 10
    date_col_start = meta_max + 1
    default_year = 2026

    # 헤더 위쪽에서 날짜 행 탐색
    date_row = None
    for r in range(max(1, header_row - 3), header_row + 1):
        hits = 0
        for c in range(date_col_start, min(ws.max_column, date_col_start + 120) + 1):
            if parse_date_header(ws.cell(r, c).value, default_year):
                hits += 1
        if hits >= 5:
            date_row = r
            break

    if date_row is None:
        date_row = header_row - 1 if header_row > 1 else header_row

    dates: list[date] = []
    c = date_col_start
    while c <= ws.max_column:
        d = parse_date_header(ws.cell(date_row, c).value, default_year)
        if d is None and dates:
            break
        if d is None:
            c += 1
            continue
        dates.append(d)
        c += 1

    if not dates:
        raise ValueError("타임라인 날짜 열을 찾을 수 없습니다")

    return dates, date_col_start


def merged_value(ws: Worksheet, row: int, col: int) -> Any:
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return ws.cell(row, col).value


def read_meta(ws: Worksheet, row: int, cols: dict[str, int], carry: dict[str, str]) -> dict[str, str]:
    out = dict(carry)
    for key, col in cols.items():
        if key == "testType":
            continue
        val = norm_text(merged_value(ws, row, col))
        if val:
            out[key] = val
    tt_col = cols.get("testType")
    if tt_col:
        out["testType"] = norm_text(merged_value(ws, row, tt_col))
    return out


def extract_bars(ws: Worksheet, row: int, dates: list[date], date_col_start: int) -> list[dict[str, Any]]:
    bars: list[dict[str, Any]] = []
    i = 0
    while i < len(dates):
        c = date_col_start + i
        cell = ws.cell(row, c)
        rgb = cell_rgb(cell)
        label = norm_text(cell.value)
        if not rgb and not label:
            i += 1
            continue
        start_i = i
        bar_type = bar_type_from_rgb(rgb)
        while i < len(dates):
            cc = date_col_start + i
            ce = ws.cell(row, cc)
            nrgb = cell_rgb(ce)
            nlabel = norm_text(ce.value)
            if not nrgb and not nlabel:
                break
            if nrgb and nrgb != rgb and i > start_i:
                break
            if nlabel and label and nlabel != label and i > start_i:
                break
            if nlabel and not label:
                label = nlabel
            if nrgb and not rgb:
                rgb = nrgb
                bar_type = bar_type_from_rgb(rgb)
            i += 1
        end_i = i - 1
        bars.append(
            {
                "start": dates[start_i].isoformat(),
                "end": dates[end_i].isoformat(),
                "type": bar_type,
                "label": label,
            }
        )
    return bars


def import_sheet(ws: Worksheet) -> tuple[list[dict[str, Any]], str | None]:
    header_row, meta_cols = find_meta_header_row(ws)
    dates, date_col_start = find_timeline(ws, header_row, meta_cols)

    rows_out: list[dict[str, Any]] = []
    carry: dict[str, str] = {
        "category": "",
        "model": "",
        "event": "",
        "variant": "",
        "manufacturer": "",
        "soc": "",
        "staff": "",
        "changes": "",
    }

    r = header_row + 1
    while r <= ws.max_row:
        meta = read_meta(ws, r, meta_cols, carry)
        tt_raw = meta.get("testType", "")
        if not meta.get("model") and not tt_raw:
            r += 1
            continue
        if meta.get("model"):
            carry = {k: meta.get(k, carry.get(k, "")) for k in carry}
        test_type = norm_test_type(tt_raw)
        if test_type not in TEST_TYPES and not tt_raw:
            r += 1
            continue

        status_col = meta_cols.get("status")
        status = norm_status(norm_text(merged_value(ws, r, status_col)) if status_col else "")
        changes_col = meta_cols.get("changes")
        changes = norm_text(merged_value(ws, r, changes_col)) if changes_col else carry.get("changes", "")

        bars: list[dict[str, Any]] = []
        if test_type == "일반성능":
            bars = extract_bars(ws, r, dates, date_col_start)

        row_id = slug_id(carry.get("model", ""), carry.get("event", ""), test_type)
        rows_out.append(
            {
                "id": f"ms-{row_id}",
                "category": carry.get("category", ""),
                "model": carry.get("model", ""),
                "event": carry.get("event", ""),
                "variant": carry.get("variant", ""),
                "manufacturer": carry.get("manufacturer", ""),
                "soc": carry.get("soc", ""),
                "staff": carry.get("staff", ""),
                "testType": test_type,
                "changes": changes,
                "status": status,
                "bars": bars,
            }
        )
        r += 1

    timeline_start = dates[0].isoformat() if dates else None
    return rows_out, timeline_start


def import_workbook(path: Path, sheet: str | None = None) -> dict[str, Any]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows, timeline_start = import_sheet(ws)
    return {"rows": rows, "count": len(rows), "timelineStart": timeline_start, "source": str(path.name)}


def main() -> None:
    ap = argparse.ArgumentParser(description="모델 검증 일정 Excel → JSON")
    ap.add_argument("xlsx", type=Path, help="입력 .xlsx")
    ap.add_argument("-o", "--output", type=Path, default=Path("scripts/seed-model-schedule-data.json"))
    ap.add_argument("--sheet", default=None, help="시트 이름 (기본: active)")
    ap.add_argument("--ts", type=Path, default=None, help="FE mock TS 출력 (optional)")
    args = ap.parse_args()

    if not args.xlsx.is_file():
        raise SystemExit(f"파일 없음: {args.xlsx}")

    payload = import_workbook(args.xlsx, args.sheet)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {args.output} — {payload['count']} rows, timelineStart={payload.get('timelineStart')}")

    if args.ts:
        ts = (
            "/** Auto-generated — scripts/import-verification-schedule-from-xlsx.py */\n"
            "import type { ModelRow } from '../utils/modelScheduleRows'\n\n"
            f"export const VERIFICATION_MOCK_ROWS: ModelRow[] = {json.dumps(payload['rows'], ensure_ascii=False, indent=2)} as ModelRow[]\n\n"
            f"export const VERIFICATION_MOCK_TIMELINE_START = {json.dumps(payload.get('timelineStart'))}\n"
        )
        args.ts.write_text(ts, encoding="utf-8")
        print(f"Wrote {args.ts}")


if __name__ == "__main__":
    main()
